from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from fpdf import FPDF

from backend.utils.text_utils import sanitize_text, split_sections

BASE_DIR = Path(__file__).resolve().parents[2]
LOGO_PATH = BASE_DIR / "assets" / "logo.png"


def _is_heading(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    if re.match(r"^(\d+\.|SECTION\s+\d+|ARTICLE\s+\d+)", stripped, re.I):
        return True
    return stripped.isupper() and len(stripped) < 100


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(text: str, doc_type: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    if LOGO_PATH.exists():
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.1))

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for block in split_sections(text):
        lines = block.splitlines()
        first = lines[0].strip() if lines else ""

        if _is_heading(first):
            p = document.add_paragraph()
            r = p.add_run(first)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            for line in lines[1:]:
                if line.strip():
                    q = document.add_paragraph(line.strip())
                    q.paragraph_format.space_after = Pt(6)
                    for rr in q.runs:
                        rr.font.name = "Times New Roman"
                        rr.font.size = Pt(11)
        else:
            for line in lines:
                if not line.strip():
                    continue
                p = document.add_paragraph(line.strip())
                p.paragraph_format.space_after = Pt(6)
                for rr in p.runs:
                    rr.font.name = "Times New Roman"
                    rr.font.size = Pt(11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("LegalEase - AI-assisted draft | Review before use")
    fr.font.name = "Times New Roman"
    fr.font.size = Pt(8)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


class BrandedPDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            self.image(str(LOGO_PATH), x=95, y=8, w=20)
            self.set_y(31)
        else:
            self.set_y(10)
        self.set_font("Helvetica", "B", 14)
        self.cell(0, 8, self.doc_type.upper(), align="C")
        self.ln(10)

    def footer(self):
        self.set_y(-14)
        self.set_font("Helvetica", size=8)
        self.cell(0, 8, "LegalEase - AI-assisted draft | Review before use", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = BrandedPDF(doc_type)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)

    for block in split_sections(text):
        lines = block.splitlines()
        for index, line in enumerate(lines):
            clean = line.strip()
            if not clean:
                continue

            if _is_heading(clean) and index == 0:
                pdf.set_font("Helvetica", "B", 11)
                pdf.multi_cell(0, 7, clean)
                pdf.ln(1)
                pdf.set_font("Helvetica", size=10)
            else:
                pdf.multi_cell(0, 6, clean)
                pdf.ln(1)
        pdf.ln(2)

    return bytes(pdf.output())
